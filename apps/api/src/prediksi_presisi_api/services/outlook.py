"""Perkiraan singkat satu bulan ke depan — di mana, jam berapa, dan apa rekomendasinya.

Permintaan pemilik proyek 9 Oktober 2026: ekspor .docx "perkiraan singkat kira-kira akan ada
ancaman di mana saja, jam berapa, dan apa rekomendasinya" untuk satu bulan ke depan.

DARI MANA PERKIRAANNYA — DINYATAKAN APA ADANYA

    Bukan model terlatih. Perkiraan bulan M disusun dari kejadian pada bulan kalender yang
    SAMA pada tahun-tahun sebelumnya yang tersedia (sampai batas tampilan data), dengan unit
    yang sama seperti rencana patroli: kelurahan × blok 3 jam × jenis. Slot diperingkat
    menurut jumlah kejadian; yang diusulkan adalah slot teratas sebanyak
    `max_slots_per_threat` dengan kejadian >= `minimum_incidents` (config/patrol/
    plan-rules.yaml — aturan yang sama, supaya dua layar tidak memakai dua aturan).

    Mengambil bulan yang sama, bukan dua belas bulan terakhir, adalah pilihan teknis:
    pola kamtibmas punya musim (mudik, tahun ajaran, akhir tahun). Pilihan ini PROPOSED,
    dan setiap slot membawa angkanya sendiri supaya dapat dihitung ulang dengan tangan.

    Rekomendasi per slot diturunkan dari config/recommendation/function-rules.yaml (fungsi
    dan tindakan per jenis ancaman) — aturan yang sama dengan rekomendasi harian.

Ini USULAN kepada Pimpinan, bukan perintah (CLAUDE.md §13–§14).
"""

from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, timedelta
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from sqlalchemy.orm import Session

from ..api.analysis import HOUR_BLOCK_STARTS, hour_block_label
from . import clock
from . import patrol_plan as planning
from .warning_issuance import load_function_rules

MONTHS_ID = (
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
)  # fmt: skip

OUTLOOK_BASIS = (
    "Perkiraan bulan sasaran disusun dari kejadian pada bulan kalender yang sama pada "
    "tahun-tahun sebelumnya yang tersedia (unit kelurahan x blok 3 jam x jenis), diperingkat "
    "menurut jumlah kejadian; slot teratas sebanyak max_slots_per_threat dengan kejadian >= "
    "minimum_incidents (config/patrol/plan-rules.yaml). Bukan model terlatih; pilihan 'bulan "
    "yang sama' adalah pilihan teknis bermusim (PROPOSED). Rekomendasi mengikuti "
    "config/recommendation/function-rules.yaml. Usulan kepada Pimpinan, bukan perintah."
)


def month_label(year: int, month: int) -> str:
    return f"{MONTHS_ID[month - 1]} {year}"


def _month_bounds(year: int, month: int) -> tuple[date, date]:
    start = date(year, month, 1)
    end = (date(year + (month // 12), month % 12 + 1, 1) - timedelta(days=1)) if True else start
    return start, end


def default_target_month() -> tuple[int, int]:
    """Bulan setelah bulan waktu acuan (aplikasi berdiri di awal 2026 → Februari 2026)."""
    now = clock.reference_now().astimezone(clock.JAKARTA)
    year, month = now.year, now.month + 1
    if month == 13:
        year, month = year + 1, 1
    return year, month


@dataclass
class OutlookSlot:
    rank: int
    polsek: str
    kecamatan: str
    kelurahan: str
    block: int
    incidents: int
    share_percent: float
    function: str
    recommendation: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "rank": self.rank,
            "polsek": self.polsek,
            "kecamatan": self.kecamatan,
            "kelurahan": self.kelurahan,
            "block_start": self.block,
            "block_label": hour_block_label(self.block),
            "incidents": self.incidents,
            "share_percent": self.share_percent,
            "function": self.function,
            "recommendation": self.recommendation,
        }


def build_outlook(
    session: Session, year: int, month: int, polsek: str | None = None
) -> dict[str, Any]:
    rules = planning.load_rules()
    function_rules = load_function_rules()
    threats = planning._threats()  # jenis yang memang dinilai versi bobot aktif
    cutoff = planning._cutoff(session)

    # Bulan yang sama pada tahun-tahun sebelumnya, selama seluruh bulannya <= batas tampilan.
    basis_months: list[tuple[int, int]] = []
    by_unit: dict[planning.Unit, int] = defaultdict(int)
    by_block: dict[tuple[str, int], int] = defaultdict(int)
    total: dict[str, int] = defaultdict(int)
    with_hour: dict[str, int] = defaultdict(int)
    for past_year in range(year - 1, year - 6, -1):
        start, end = _month_bounds(past_year, month)
        if end > cutoff:
            continue
        observed = planning.observe(session, start, end, threats, polsek)
        if not observed.total:
            continue
        basis_months.append((past_year, month))
        for unit, count in observed.by_unit.items():
            by_unit[unit] += count
        for key, count in observed.by_block.items():
            by_block[key] += count
        for threat, count in observed.total.items():
            total[threat] += count
            with_hour[threat] += observed.with_hour.get(threat, 0)

    per_threat: list[dict[str, Any]] = []
    for threat in threats:
        rule = function_rules.for_threat(threat)
        candidates = sorted(
            ((unit, count) for unit, count in by_unit.items() if unit.threat_type == threat),
            key=lambda item: (-item[1], item[0].kelurahan, item[0].block),
        )
        known = with_hour.get(threat, 0)
        slots: list[OutlookSlot] = []
        for unit, count in candidates:
            if count < rules.minimum_incidents or len(slots) >= rules.max_slots_per_threat:
                break
            label = hour_block_label(unit.block)
            slots.append(
                OutlookSlot(
                    rank=len(slots) + 1,
                    polsek=unit.polsek,
                    kecamatan=unit.kecamatan,
                    kelurahan=unit.kelurahan,
                    block=unit.block,
                    incidents=count,
                    share_percent=round(100 * count / known, 1) if known else 0.0,
                    function=rule.function,
                    recommendation=(
                        f"{rule.action} oleh {rule.function} di Kelurahan {unit.kelurahan} "
                        f"({unit.kecamatan}) pada {label}."
                    ),
                )
            )
        blocks: list[dict[str, Any]] = [
            {
                "block_start": start,
                "block_label": hour_block_label(start),
                "incidents": by_block.get((threat, start), 0),
                "share_percent": (
                    round(100 * by_block.get((threat, start), 0) / known, 1) if known else 0.0
                ),
            }
            for start in HOUR_BLOCK_STARTS
        ]
        peak = max(blocks, key=lambda row: int(row["incidents"])) if known else None
        per_threat.append(
            {
                "threat_type": threat,
                "function": rule.function,
                "action": rule.action,
                "basis_incidents": total.get(threat, 0),
                "basis_with_hour": known,
                "peak_block": peak["block_label"] if peak and peak["incidents"] else None,
                "hour_profile": blocks,
                "slots": [slot.as_dict() for slot in slots],
            }
        )

    return {
        "target_month": f"{year:04d}-{month:02d}",
        "target_label": month_label(year, month),
        "basis_months": [f"{y:04d}-{m:02d}" for y, m in basis_months],
        "basis_labels": [month_label(y, m) for y, m in basis_months],
        "scope": polsek,
        "rules_version": rules.version,
        "status": "PROPOSED",
        "threats": per_threat,
        "basis": OUTLOOK_BASIS,
        "function_rules_status": function_rules.status,
    }


# ---------------------------------------------------------------------------
# .docx
# ---------------------------------------------------------------------------


ROMAN = ("I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII")
FONT = "Arial"
#: Klasifikasi pada kaki halaman — mengikuti contoh Kirkat pemilik proyek (13 Oktober 2026).
#: Satu tempat; ganti di sini bila kebijakan klasifikasi menetapkan lain.
CLASSIFICATION = "RAHASIA"


def _run(paragraph: Any, text: str, *, bold: bool = False, size: int = 12) -> Any:
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)
    return run


def _set_cell(cell: Any, text: str, *, bold: bool = False, center: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _run(paragraph, text, bold=bold, size=10)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def render_docx(outlook: dict[str, Any]) -> bytes:
    """Dokumen Word mengikuti bentuk **Perkiraan Intelijen Singkat (Kirkat)** Polri, dari contoh
    yang diberikan pemilik proyek 13 Oktober 2026: A4, Arial 12, margin kiri 2,5 cm dan 2 cm
    lainnya, kop satuan dan nomor naskah, judul "PERKIRAAN INTELIJEN SINGKAT / TENTANG / …",
    bagian PENDAHULUAN – KEADAAN SASARAN – ANALISA – KESIMPULAN – SARAN dengan penomoran
    bertingkat dan teks rata kiri-kanan, blok tanda tangan, serta klasifikasi di kaki halaman.
    Isinya dari data (`build_outlook`); kalimat-kalimatnya aturan, bukan model bahasa."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = section.top_margin = section.bottom_margin = Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(footer, CLASSIFICATION, bold=True)

    def para(
        text: str = "",
        *,
        bold: bool = False,
        align: str = "justify",
        indent_cm: float = 0.0,
        hanging_cm: float = 0.0,
        after: int = 6,
    ) -> Any:
        paragraph = document.add_paragraph()
        if text:
            _run(paragraph, text, bold=bold)
        paragraph.alignment = {
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }.get(align, WD_ALIGN_PARAGRAPH.JUSTIFY)
        fmt = paragraph.paragraph_format
        fmt.left_indent = Cm(indent_cm + hanging_cm)
        fmt.first_line_indent = Cm(-hanging_cm) if hanging_cm else None
        fmt.space_after = Pt(after)
        return paragraph

    def heading(number: str, title: str) -> None:
        para(f"{number}\t{title}", bold=True, align="left", hanging_cm=1.0, after=6)

    def numbered(label: str, text: str, level: int = 1, *, bold: bool = False) -> None:
        indent = 1.0 * (level - 1)
        para(f"{label}\t{text}", bold=bold, indent_cm=indent, hanging_cm=1.0, after=4)

    now = clock.reference_now().astimezone(clock.JAKARTA)
    scope = outlook.get("scope")
    target = outlook["target_label"]
    target_upper = target.upper()
    basis = ", ".join(outlook["basis_labels"]) or "tidak ada bulan dasar yang tersedia"
    threats = outlook["threats"]
    names = ", ".join(t["threat_type"] for t in threats) or "—"
    month_roman = ROMAN[int(outlook["target_month"][5:7]) - 1]

    # ---- Kop dan nomor ----
    para("KEPOLISIAN NEGARA REPUBLIK INDONESIA", align="center", after=0)
    para("DAERAH METRO JAYA", align="center", after=0)
    para(
        f"RESOR METRO JAKARTA SELATAN{(' — SEKTOR ' + scope.upper()) if scope else ''}",
        align="center",
        after=0,
    )
    para(f"Nomor: R/Kirkat -        /{month_roman}/{now.year}/________", align="center", after=18)

    # ---- Judul ----
    para("PERKIRAAN INTELIJEN SINGKAT", bold=True, align="center", after=6)
    para("TENTANG", bold=True, align="center", after=6)
    para(
        f"PERKIRAAN KERAWANAN {names} DI WILAYAH HUKUM POLRES METRO JAKARTA SELATAN",
        bold=True,
        align="center",
        after=0,
    )
    para(f"PADA BULAN {target_upper}", bold=True, align="center", after=18)

    # ---- I. PENDAHULUAN ----
    heading("I.", "PENDAHULUAN")
    numbered(
        "1.",
        f"Kejadian {names} merupakan gangguan kamtibmas yang berulang di wilayah hukum Polres "
        f"Metro Jakarta Selatan. Pola tempat dan waktu kejadiannya pada bulan-bulan sebelumnya "
        f"dapat menjadi dasar perkiraan kerawanan bulan {target}.",
    )
    numbered(
        "2.",
        "Perkiraan Intelijen Singkat ini disampaikan kepada Pimpinan sebagai masukan untuk "
        "dijadikan bahan pertimbangan dalam menentukan sasaran, waktu, dan bentuk kegiatan "
        f"kepolisian pada bulan {target}. Perkiraan disusun oleh sistem KRESNA SIGAP (Sistem "
        "Inteligensi Geospasial dan Analitik Prediktif) dari data kejadian yang tercatat, "
        f"dengan status {outlook['status']}: usulan, bukan perintah.",
    )
    para("", after=6)

    # ---- II. KEADAAN SASARAN ----
    heading("II.", "KEADAAN SASARAN")
    numbered("1.", "Subyek", bold=True)
    para(
        f"Kerawanan {names} pada bulan {target} di seluruh kelurahan wilayah hukum Polres Metro "
        f"Jakarta Selatan{(' (Polsek ' + scope + ')') if scope else ''}.",
        indent_cm=1.0,
        after=6,
    )
    numbered("2.", "Situasi Khusus", bold=True)
    numbered(
        "a.",
        f"Dasar perkiraan adalah kejadian pada bulan kalender yang sama tahun-tahun sebelumnya "
        f"({basis}), dengan satuan analisis kelurahan, blok waktu tiga jam, dan jenis kejahatan.",
        level=2,
    )
    for letter, threat in zip("bcdefgh", threats, strict=False):
        peak = threat["peak_block"] or "tidak dapat ditentukan"
        numbered(
            f"{letter}.",
            f"{threat['threat_type']}: {threat['basis_incidents']} kejadian pada bulan dasar "
            f"({threat['basis_with_hour']} dengan jam tercatat); blok waktu paling rawan {peak}.",
            level=2,
        )
    numbered("3.", "Data Sasaran", bold=True)
    for letter, threat in zip("abcdefgh", threats, strict=False):
        numbered(
            f"{letter}.",
            f"{threat['threat_type']} — kelurahan dan blok waktu dengan kejadian berulang:",
            level=2,
        )
        if not threat["slots"]:
            para(
                "Tidak ada kelurahan dan blok waktu yang memenuhi batas minimum kejadian; "
                "polanya tersebar.",
                indent_cm=2.0,
                after=6,
            )
            continue
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = (Cm(1.2), Cm(7.0), Cm(3.2), Cm(3.4))
        for cell, text, width in zip(
            table.rows[0].cells,
            ("NO", "KELURAHAN (KECAMATAN)", "WAKTU", "KEJADIAN"),
            widths,
            strict=True,
        ):
            _set_cell(cell, text, bold=True, center=True)
            cell.width = width
        for slot in threat["slots"]:
            row = table.add_row().cells
            values = (
                str(slot["rank"]),
                f"{slot['kelurahan']} ({slot['kecamatan']})",
                slot["block_label"],
                f"{slot['incidents']} ({slot['share_percent']}%)",
            )
            for index, (cell, text, width) in enumerate(zip(row, values, widths, strict=True)):
                _set_cell(cell, text, center=index != 1)
                cell.width = width
        para("", after=6)

    # ---- III. ANALISA ----
    heading("III.", "ANALISA")
    numbered("1.", "Kemampuan Sasaran (Ancaman)", bold=True)
    for letter, threat in zip("abcdefgh", threats, strict=False):
        top = threat["slots"][:3]
        where = (
            "; ".join(
                f"{s['kelurahan']} pada {s['block_label']} ({s['incidents']} kejadian)" for s in top
            )
            if top
            else "tidak terpusat pada kelurahan tertentu"
        )
        share = sum(s["share_percent"] for s in threat["slots"])
        numbered(
            f"{letter}.",
            f"{threat['threat_type']} diperkirakan paling rawan pada blok waktu "
            f"{threat['peak_block'] or 'yang belum dapat ditentukan'}; kelurahan dengan pola "
            f"berulang terkuat: {where}. Seluruh slot usulan mencakup {round(share, 1)}% kejadian "
            "berjam pada bulan dasar.",
            level=2,
        )
    numbered("2.", "Kemampuan Penanggulangan", bold=True)
    for letter, threat in zip("abcdefgh", threats, strict=False):
        numbered(
            f"{letter}.",
            f"{threat['threat_type']}: {threat['action']} oleh fungsi {threat['function']} pada "
            "kelurahan dan blok waktu yang disebut di atas, dengan dukungan fungsi Binmas untuk "
            "imbauan kepada masyarakat.",
            level=2,
        )
    para("", after=6)

    # ---- IV. KESIMPULAN ----
    heading("IV.", "KESIMPULAN")
    numbered("1.", "Kondisi Sasaran", bold=True)
    para(
        f"Kerawanan {names} bulan {target} diperkirakan mengikuti pola bulan yang sama pada "
        "tahun-tahun sebelumnya: terpusat pada kelurahan dan blok waktu tertentu yang tercantum "
        "pada Data Sasaran. Angka ini dihitung dari kejadian yang tercatat, bukan dari model "
        "terlatih, dan belum dievaluasi terhadap bulan sasaran.",
        indent_cm=1.0,
        after=6,
    )
    numbered("2.", "Kemampuan Satuan", bold=True)
    para(
        "Satuan memiliki kemampuan preemtif, preventif, dan represif untuk menekan kerawanan "
        "tersebut melalui penempatan patroli pada kelurahan dan blok waktu rawan, penjagaan "
        "objek vital, dan imbauan kepada masyarakat.",
        indent_cm=1.0,
        after=6,
    )

    # ---- V. SARAN ----
    heading("V.", "SARAN")
    number = 1
    for threat in threats:
        for slot in threat["slots"][:5]:
            numbered(
                f"{number}.",
                f"{threat['action']} oleh {slot['function']} di Kelurahan {slot['kelurahan']} "
                f"({slot['kecamatan']}) pada {slot['block_label']} — {threat['threat_type']}.",
            )
            number += 1
        if not threat["slots"] and threat["peak_block"]:
            numbered(
                f"{number}.",
                f"{threat['action']} oleh {threat['function']} pada blok waktu "
                f"{threat['peak_block']} di seluruh wilayah — {threat['threat_type']}.",
            )
            number += 1
    numbered(
        f"{number}.",
        "Penetapan kegiatan dan pengerahan kekuatan tetap diputuskan Pimpinan; keputusan atas "
        "usulan ini tercatat pada jejak audit sistem untuk dievaluasi pada akhir bulan.",
    )
    para("", after=18)

    # ---- Penutup / tanda tangan ----
    para(f"Jakarta,      {MONTHS_ID[now.month - 1]} {now.year}", align="center", after=12)
    para("KEPALA KEPOLISIAN RESOR METRO JAKARTA SELATAN", align="right", after=48)
    para("________________________", align="right", after=12)
    para("Autentikasi: ______________", align="left", after=0)
    para("Paraf: Kasat Intelkam: ______   Kabag Ops: ______", align="left", after=0)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
